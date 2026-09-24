from pathlib import Path

from docx import Document

from backend.services.resume_intelligence import (
    analyze_resume_file,
    analyze_resume_text,
    clean_resume_text,
    cosine_similarity,
)


def test_resume_text_analysis_extracts_sections_and_skills():
    result = analyze_resume_text(
        """SUMMARY\nBackend engineer\n\nSKILLS\nPython, FastAPI, PostgreSQL\n\nEXPERIENCE\nBuilt APIs\n\nEDUCATION\nBS Computer Science\n\nPROJECTS\nCareerLens"""
    )

    assert result['overall_score'] > 0
    assert result['keyword_score'] == 100.0
    assert {'Python', 'FastAPI', 'PostgreSQL'} <= set(result['extracted_skills'])
    assert result['missing_keywords'] == []
    assert result['section_analysis']['experience']['present'] is True


def test_docx_extraction_and_text_cleaning(tmp_path: Path):
    document = Document()
    document.add_paragraph('  SKILLS  ')
    document.add_paragraph('Python   FastAPI')
    path = tmp_path / 'resume.docx'
    document.save(path)

    result = analyze_resume_file(str(path), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    assert result['extracted_text'] == 'SKILLS\nPython FastAPI'
    assert result['extracted_skills'] == ['Python', 'FastAPI']
    assert clean_resume_text('A\x00\n\n\nB') == 'A\n\nB'


def test_cosine_similarity_handles_normalized_vectors_and_invalid_shapes():
    assert cosine_similarity([1.0, 0.0], [1.0, 0.0]) == 1.0
    assert cosine_similarity([1.0], [1.0, 0.0]) == 0.0
    assert cosine_similarity([0.0, 0.0], [1.0, 0.0]) == 0.0
